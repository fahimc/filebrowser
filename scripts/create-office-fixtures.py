"""Create disposable Office samples for scripts/office-smoke.mjs.

Requires python-docx, openpyxl, python-pptx and reportlab (test tools only).
Existing files are left intact; pass a new empty folder for a clean run.
"""
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from reportlab.pdfgen import canvas

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)

if not (root / "Quarterly report.docx").exists():
    doc = Document()
    doc.add_heading("Quarterly report", 0)
    doc.add_paragraph("A real Word document stored in a normal Windows folder.")
    doc.add_heading("Project progress", 1)
    doc.add_paragraph("Browser editing and automatic saving are ready for review.")
    doc.save(root / "Quarterly report.docx")

if not (root / "Team budget.xlsx").exists():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    for row in [["Department", "Budget"], ["Engineering", 42000], ["Design", 18000], ["Operations", 25000], ["Total", "=SUM(B2:B4)"]]:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 18
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="4285F4")
    workbook.save(root / "Team budget.xlsx")

if not (root / "Team update.pptx").exists():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Team update"
    slide.placeholders[1].text = "FileBrowser Drive + browser Office editor"
    presentation.save(root / "Team update.pptx")

if not (root / "Annual report.pdf").exists():
    pdf = canvas.Canvas(str(root / "Annual report.pdf"))
    pdf.setTitle("Annual report")
    pdf.setFont("Helvetica-Bold", 24)
    pdf.drawString(72, 750, "Annual report 2026")
    pdf.setFont("Helvetica", 12)
    pdf.drawString(72, 714, "This ordinary PDF can be viewed and annotated in the browser.")
    pdf.drawString(72, 690, "Edits are saved back to this Windows folder.")
    pdf.showPage()
    pdf.save()

print(f"Office samples are ready in {root.resolve()}")
